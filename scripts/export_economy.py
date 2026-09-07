#!/usr/bin/env python3
"""Export TechWorks xlsx → public/economy.json for the projector board.

Usage:
  python3 scripts/export_economy.py [path/to/TechWorks_4.0_SalarySystem.xlsx]
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.stderr.write("pip install openpyxl\n")
    raise

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "artifacts" / "TechWorks_4.0_SalarySystem.xlsx"
OUT = ROOT / "public" / "economy.json"
CODES = {"3": 25, "2": 20, "1": 15, "A": 0, "E": 0, "P": -25}


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return v


def main() -> None:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = load_workbook(src, data_only=False)
    cfg = wb["CONFIG"]
    keys = {str(cfg.cell(r, 1).value): cfg.cell(r, 2).value for r in range(2, 50) if cfg.cell(r, 1).value}
    bell = []
    for r in range(3, 9):
        p, g = cfg.cell(r, 5).value, cfg.cell(r, 6).value
        if p is None:
            continue
        try:
            bell.append({"period": int(p), "grade": int(g or 0)})
        except (TypeError, ValueError):
            continue
    meta = {
        "title": "TechWorks",
        "quarterName": str(keys.get("SessionName") or keys.get("QuarterName") or "S1"),
        "currentWeek": int(keys.get("CurrentWeek") or 1),
        "schoolYear": str(keys.get("SchoolYear") or "2026-27"),
        "codes": {
            "3": int(keys.get("Code3") or 25),
            "2": int(keys.get("Code2") or 20),
            "1": int(keys.get("Code1") or 15),
            "A": int(keys.get("CodeA") or 0),
            "E": int(keys.get("CodeE") or 0),
            "P": int(keys.get("CodeP") or -25),
        },
        "bell": bell
        or [
            {"period": 1, "grade": 6},
            {"period": 2, "grade": 8},
            {"period": 3, "grade": 7},
            {"period": 8, "grade": 7},
            {"period": 9, "grade": 8},
            {"period": 10, "grade": 6},
        ],
        "market": {
            "source": "DOW" if str(keys.get("MarketSource") or "TEACHER").upper() == "DOW" else "TEACHER",
            "index": float(keys.get("MarketManual") or 45000),
            "baseline": float(keys.get("MarketBaseline") or 45000),
            "shock": float(keys.get("MarketShock") or 0),
            "factor": (float(keys.get("MarketManual") or 45000) / float(keys.get("MarketBaseline") or 45000))
            * (1 + float(keys.get("MarketShock") or 0) / 100)
            if float(keys.get("MarketBaseline") or 0)
            else 1,
        },
    }
    crews = []
    cr = wb["CREWS"]
    for r in range(3, 40):
        per = cr.cell(r, 1).value
        if per is None:
            continue
        try:
            period = int(per)
        except (TypeError, ValueError):
            continue
        crews.append({"period": period, "key": str(cr.cell(r, 2).value), "name": str(cr.cell(r, 3).value)})

    ros = wb["ROSTER"]
    kul = wb["KULIBERT ONLY"]
    students = []
    for i in range(80):
        rr = 3 + i
        kr = 4 + i
        sid = ros.cell(rr, 1).value
        if not sid:
            continue
        days = []
        for col in (5, 6, 7, 8):
            v = kul.cell(kr, col).value
            days.append("" if v is None else str(v))
        opening = 0
        if "BANK" in wb.sheetnames:
            try:
                opening = float(wb["BANK"].cell(kr, 6).value or 0)
            except (TypeError, ValueError):
                opening = 0
        students.append(
            {
                "id": str(sid),
                "first": str(ros.cell(rr, 2).value or ""),
                "last": str(ros.cell(rr, 3).value or ""),
                "period": int(ros.cell(rr, 5).value or 0),
                "grade": int(ros.cell(rr, 9).value or 0) if str(ros.cell(rr, 9).value or "").isdigit() else None,
                "crewKey": str(ros.cell(rr, 6).value or "Crew A"),
                "days": days,
                "bonus": float(kul.cell(kr, 10).value or 0),
                "deduct": float(kul.cell(kr, 11).value or 0),
                "clutch": float(kul.cell(kr, 12).value or 0),
                "opening": opening,
            }
        )

    OUT.write_text(json.dumps({"meta": meta, "crews": crews, "students": students}, indent=2) + "\n")
    src_copy = ROOT / "src" / "data" / "economy.json"
    src_copy.parent.mkdir(parents=True, exist_ok=True)
    src_copy.write_text(OUT.read_text())
    print(f"Wrote {OUT} and {src_copy} ({len(students)} students)")


if __name__ == "__main__":
    main()
