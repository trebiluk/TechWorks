#!/usr/bin/env python3
"""TechWorks default 16:9 navy + violet deck."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt
from lxml import etree

NAVY = RGBColor(0x05, 0x08, 0x16)
SURFACE = RGBColor(0x14, 0x1C, 0x42)
VIOLET = RGBColor(0x8B, 0x6C, 0xFF)
ROYAL = RGBColor(0x1E, 0x4B, 0xAF)
CYAN = RGBColor(0x2E, 0xE6, 0xFF)
GOLD = RGBColor(0xF0, 0xD4, 0x8A)
PAPER = RGBColor(0xF7, 0xF9, 0xFF)
MUTED = RGBColor(0xB7, 0xC4, 0xEA)

W, H = Inches(13.333), Inches(7.5)
ROOT = Path("/workspace")
MARK = ROOT / "public/brand/techworks.png"
BERT = ROOT / "public/berty/brand/bertybot_waving.svg"
OUT_A = ROOT / "artifacts/TechWorks-Deck.pptx"
OUT_P = ROOT / "public/TechWorks-Deck.pptx"
FOOT = "© 2026 Richard Kulibert. TECHWORKS™ v1.83.1."


def rgb(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def box(s, l, t, w, h, color):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    rgb(sh, color)
    return sh


def roundbox(s, l, t, w, h, color):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    rgb(sh, color)
    return sh


def txt(s, l, t, w, h, text, size=18, color=PAPER, bold=False, font="Calibri", align=PP_ALIGN.LEFT):
    sh = s.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return sh


def plate(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, NAVY)
    box(s, 0, 0, Inches(0.18), H, VIOLET)
    return s


def kicker(s, text):
    txt(s, Inches(0.55), Inches(0.28), Inches(8), Inches(0.35), text.upper(), 11, GOLD, True, "Consolas")


def footer(s):
    txt(s, Inches(0.55), Inches(7.12), Inches(10), Inches(0.28), FOOT, 10, MUTED, False, "Consolas")
    if MARK.exists():
        s.shapes.add_picture(str(MARK), Inches(10.55), Inches(0.22), height=Inches(0.38))


def cards(s, items, cols=2, y=Inches(2.05)):
    n = len(items)
    cols = min(cols, n)
    rows = (n + cols - 1) // cols
    gap = Inches(0.18)
    left = Inches(0.55)
    usable = W - Inches(0.9)
    cw = (usable - gap * (cols - 1)) / cols
    ch = Inches(1.85) if rows == 1 else Inches(2.05)
    for i, item in enumerate(items):
        r, c = divmod(i, cols)
        x = left + c * (cw + gap)
        yy = y + r * (ch + gap)
        roundbox(s, x, yy, cw, ch, SURFACE)
        title, line, num = item[0], item[1], item[2] if len(item) > 2 else ""
        if num:
            txt(s, x + Inches(0.22), yy + Inches(0.12), cw - Inches(0.4), Inches(0.55), num, 28, CYAN, True, "Consolas")
            txt(s, x + Inches(0.22), yy + Inches(0.7), cw - Inches(0.4), Inches(0.35), title, 16, GOLD, True, "Consolas")
            txt(s, x + Inches(0.22), yy + Inches(1.05), cw - Inches(0.4), Inches(0.75), line, 16, PAPER)
        else:
            txt(s, x + Inches(0.28), yy + Inches(0.28), cw - Inches(0.5), Inches(0.35), title.upper(), 12, GOLD, True, "Consolas")
            txt(s, x + Inches(0.28), yy + Inches(0.7), cw - Inches(0.5), Inches(1.05), line, 18, PAPER)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    s = plate(prs)
    kicker(s, "Solvay Tech Ed  ·  Room 13  ·  2026–27")
    txt(s, Inches(0.55), Inches(2.1), Inches(12), Inches(1.8), "TECHWORKS", 80, CYAN, True)
    txt(s, Inches(0.55), Inches(4.05), Inches(11), Inches(0.7), "Workshop you can see.", 28, MUTED)
    txt(s, Inches(0.55), Inches(4.85), Inches(11), Inches(0.4), "Aliases on the wall. Legal names stay in the vault.", 14, GOLD, False, "Consolas")
    footer(s)

    s = plate(prs)
    kicker(s, "How this class works")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.8), "Crew. Skills. Gold.", 36, PAPER, True)
    cards(s, [
        ("Crew", "You work in a small crew. Four people, one bench."),
        ("Skills", "Measure, cut, finish, share. Real work you can see."),
        ("Gold XP", "Getting better at the craft. That is the point."),
        ("Cash", "A perk game. Not the grade. Not on the family sheet."),
    ])
    footer(s)

    s = plate(prs)
    kicker(s, "Every period")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "Four beats", 36, PAPER, True)
    cards(s, [
        ("ENTER", "Sit with your crew.", "1"),
        ("LISTEN", "Directions first. Then questions.", "2"),
        ("CREW WORK", "Today’s activity. Tools with a purpose.", "3"),
        ("CLEAN UP", "Stations reset before the bell.", "4"),
    ], cols=4, y=Inches(1.7))
    footer(s)

    s = plate(prs)
    kicker(s, "Today’s effort  ·  not the skill grade")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "3  ·  2  ·  1", 36, PAPER, True)
    cards(s, [
        ("On the job", "Full day with the crew.", "3"),
        ("Needs a nudge", "Working, with a check-in.", "2"),
        ("Not with the crew", "Off-task. Redirect.", "1"),
    ], cols=3, y=Inches(1.8))
    txt(s, Inches(0.55), Inches(6.55), Inches(12), Inches(0.4), "A absent · E excused · P present but no work. Crew lead taps this period only.", 12, MUTED)
    footer(s)

    s = plate(prs)
    kicker(s, "Observable  ·  Watch stores the sentence")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "Skills 1 to 4", 36, PAPER, True)
    cards(s, [
        ("Beginning", "Needs a demo. Not independent yet.", "1"),
        ("Developing", "Can do it with a check-in.", "2"),
        ("Proficient", "Independent. Meets the standard.", "3"),
        ("Distinguished", "Can teach a crewmate. Exceeds.", "4"),
    ], cols=4, y=Inches(1.7))
    txt(s, Inches(0.55), Inches(6.5), Inches(12), Inches(0.4), "SAFETY · MEASURE · DRAW · MODEL · TOOLS · FINISH · PRESENT · TEAM", 12, GOLD, False, "Consolas")
    footer(s)

    s = plate(prs)
    kicker(s, "On the unit  ·  not a second score")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "S  T  E  M", 36, PAPER, True)
    cards(s, [
        ("Science", "Materials, force, speed, and what the test showed.", "S"),
        ("Technology", "Tools, files, and the process that made the part.", "T"),
        ("Engineering", "The design: constraints, ideas, and the next change.", "E"),
        ("Math", "Measure, size, scale, and whether the numbers hold.", "M"),
    ], cols=4, y=Inches(1.7))
    txt(s, Inches(0.55), Inches(6.5), Inches(12), Inches(0.4), "Grade 6 · How can a small force move a bigger load?", 14, MUTED)
    footer(s)

    s = plate(prs)
    kicker(s, "No work until this is solid")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "Safety first", 36, PAPER, True)
    cards(s, [
        ("Glasses", "On before the tool starts."),
        ("Ask", "Ask, then wait for the nod."),
        ("Zone", "Stand clear of swing and offcut."),
        ("License", "One tool at a time. Licensed tools only."),
    ])
    txt(s, Inches(0.55), Inches(6.55), Inches(12), Inches(0.35), "Distinguished: stop a crewmate who skipped PPE.", 13, GOLD)
    footer(s)

    s = plate(prs)
    kicker(s, "Duplicate this slide every period")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "Do this now", 36, CYAN, True)
    rows = [
        ("ACTIVITY", "Technical Drawing", "Today’s job at the bench."),
        ("LOOK FOR", "a 3", "Independent. Meets the standard."),
        ("QUESTION", "How can a small force move a bigger load?", "Driving question for this unit."),
    ]
    y = Inches(1.7)
    for k, title, line in rows:
        roundbox(s, Inches(0.55), y, Inches(12.2), Inches(1.35), SURFACE)
        txt(s, Inches(0.8), y + Inches(0.22), Inches(2.2), Inches(0.9), k, 12, GOLD, True, "Consolas")
        txt(s, Inches(3.1), y + Inches(0.18), Inches(9.2), Inches(0.55), title, 24, PAPER, True)
        txt(s, Inches(3.1), y + Inches(0.72), Inches(9.2), Inches(0.4), line, 14, MUTED)
        y += Inches(1.55)
    footer(s)

    s = plate(prs)
    kicker(s, "Last beat  ·  before the bell")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "Clean up", 36, PAPER, True)
    cards(s, [
        ("Tools", "Tools and kits away."),
        ("Bench", "Floor and tables clear."),
        ("Seats", "Names stay until the room is ready."),
        ("Help", "Caught helping extra can earn a perk — not XP."),
    ])
    footer(s)

    s = plate(prs)
    kicker(s, "Section")
    txt(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.7), "Title here", 36, PAPER, True)
    cards(s, [
        ("Left", "Type over this. Keep the navy plate."),
        ("Right", "Duplicate for demos, critiques, club."),
    ], cols=2, y=Inches(1.9))
    txt(s, Inches(0.55), Inches(6.5), Inches(12), Inches(0.4), "Default theme master. Do not change the colors.", 13, GOLD)
    footer(s)

    s = plate(prs)
    kicker(s, "Aliases on the projector")
    txt(s, Inches(0.55), Inches(2.3), Inches(12), Inches(1.4), "See you at the bench.", 44, PAPER, True)
    txt(s, Inches(0.55), Inches(3.85), Inches(12), Inches(0.5), "Workshop, not a lecture.", 22, MUTED)
    txt(s, Inches(0.55), Inches(4.55), Inches(12), Inches(0.4), "Legal names stay in the vault.", 14, GOLD, False, "Consolas")
    footer(s)

    OUT_A.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_A))
    OUT_P.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_P))
    print(f"wrote {OUT_A} and {OUT_P} ({OUT_A.stat().st_size} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
